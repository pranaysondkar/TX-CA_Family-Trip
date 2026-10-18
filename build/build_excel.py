"""Builds ../../Family_Trip_Dec2026_Plan.xlsx from trip_data.py.

Run:  python build_excel.py   (needs openpyxl)
"""
import datetime as dt
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import trip_data as T


def live_bookings():
    """Current Bookings rows from the online sheet (id -> (booked, actual)). Empty if not reachable."""
    import subprocess
    if not T.BOOKINGS_API:
        return {}
    try:
        raw = subprocess.run(["curl", "-sL", "-m", "30", T.BOOKINGS_API], capture_output=True, check=True).stdout
        rows = json.loads(raw)["rows"]
    except Exception:
        return {}
    return {r["id"]: (bool(r.get("booked")), r.get("actual")) for r in rows}


LIVE = live_bookings()
print("live bookings rows:", len(LIVE))

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "..", "Family_Trip_Dec2026_Plan.xlsx")

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3A5F")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color="0000FF")
BOLD = Font(name=F, size=10, bold=True)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
GREEN = Font(name=F, size=10, color="008000")
TITLE = Font(name=F, size=16, bold=True, color="1F3A5F")
SUB = Font(name=F, size=10, italic=True, color="555555")
YELLOW = PatternFill("solid", fgColor="FFF59D")
TX_FILL = PatternFill("solid", fgColor="FDF1E7")
CA_FILL = PatternFill("solid", fgColor="E8F1FB")
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
WRAP = Alignment(wrap_text=True, vertical="top")
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;-0.0%;"-"'


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = H_FONT, H_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def link(cell, key, text=None):
    url = T.L.get(key) if key else None
    if url:
        cell.value = text or "Open ↗"
        cell.hyperlink = url
        cell.font = LINK


wb = Workbook()
BOOK_ROWS = []

# ------------------------------------------------------------------ Budget (built first; Summary references it)
bs = wb.active
bs.title = "Budget"
cols = ["Category", "Item", "Date(s)", "Qty", "Unit cost ($)", "Subtotal ($)", "Tax / fees %", "In plan? (Y/N)",
        "Total ($)", "Status", "Lean plan ($)", "Book / info", "Notes & source", "Actual paid ($)", "Projected ($)", "ID"]
header(bs, 1, cols, [20, 52, 18, 7, 12, 12, 10, 9, 12, 13, 13, 12, 70, 13, 13, 30])
r = 2
for (cat, item, dates, qty, unit, tax, status, lean, lk, note) in T.BUDGET:
    bs.cell(r, 1, cat); bs.cell(r, 2, item); bs.cell(r, 3, dates)
    q = bs.cell(r, 4, qty); u = bs.cell(r, 5, unit); t = bs.cell(r, 7, tax)
    for c in (q, u, t):
        c.font = BLUE
    bs.cell(r, 6, f"=D{r}*E{r}").number_format = USD
    inc = bs.cell(r, 8, "N" if status == "Optional" else "Y"); inc.font = BLUE
    inc.alignment = Alignment(horizontal="center")
    bs.cell(r, 9, f'=IF(H{r}="N",0,F{r}*(1+G{r}))').number_format = USD
    bid = T.budget_id(cat, item)
    bs.cell(r, 16, bid)
    look = lambda col: f"INDEX(Bookings!${col}:${col},MATCH($P{r},Bookings!$A:$A,0))"
    bs.cell(r, 10, f'=IFERROR(IF({look("G")}="Yes","Booked","{status}"),"{status}")')
    lc = bs.cell(r, 11, f"=I{r}" if lean is None else round(lean, 2))
    lc.number_format = USD
    if lean is not None:
        lc.font = BLUE
    link(bs.cell(r, 12), lk)
    bs.cell(r, 13, note)
    ac = bs.cell(r, 14, f'=IFERROR(IF({look("H")}="","",{look("H")}),"")'); ac.font = GREEN; ac.number_format = USD2
    bs.cell(r, 15, f'=IF(N{r}="",I{r},N{r})').number_format = USD
    booked, actual = LIVE.get(bid, (False, None))
    BOOK_ROWS.append((bid, cat, item, dates, "N" if status == "Optional" else "Y", round(qty * unit * (1 + tax), 2),
                      "Yes" if booked or cat == "Already paid" or status == "Booked" else "No",
                      actual if actual is not None else (round(qty * unit * (1 + tax), 2) if cat == "Already paid" else None)))
    if booked or actual is not None or status == "Booked":
        status = "Booked"
    u.number_format = USD2
    t.number_format = PCT
    for c in range(1, 17):
        cell = bs.cell(r, c)
        if cell.font not in (BLUE, GREEN) and c != 12:
            cell.font = BODY
        cell.alignment = WRAP
        cell.border = THIN
    if status in ("To book", "Reserve", "Book", "Book early", "Confirm", "Decide", "Get quote"):
        bs.cell(r, 10).fill = YELLOW
    r += 1
last = r - 1
BUDGET_TOTAL_ROW = r
bs.cell(r, 2, "TOTAL").font = BOLD
for col, letter in ((9, "I"), (11, "K"), (14, "N"), (15, "O")):
    c = bs.cell(r, col, f"=SUM({letter}2:{letter}{last})")
    c.font, c.number_format = BOLD, USD
tab = Table(displayName="BudgetLines", ref=f"A1:P{last}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
bs.add_table(tab)
bs.cell(r + 2, 1, "Legend:").font = BOLD
bs.cell(r + 2, 2, "Blue = inputs you can change (qty, price, tax, In plan Y/N). Black = formulas. Green = pulled from the Bookings tab. "
        "Mark bookings and actual costs on the website or in the Bookings tab, not here. Optional items start as N. "
        "Yellow status = still needs booking or a decision. Lean plan column = cheaper version "
        "(0 = dropped). Prices checked live " + T.PRICES_CHECKED + ".").font = SUB
bs.cell(r + 2, 2).alignment = WRAP

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary", 0)
sm.column_dimensions["A"].width = 34
for c in "BCD":
    sm.column_dimensions[c].width = 16
sm.column_dimensions["E"].width = 60
sm["A1"] = "Family Trip · Dec 6, 2026 – Jan 4, 2027"; sm["A1"].font = TITLE
sm["A2"] = f"{T.TRAVELERS}. Texas visit + California road trip Dec 19–27. Prices checked live {T.PRICES_CHECKED}."
sm["A2"].font = SUB
row = 4
for i, h in enumerate(["Category", "Planned ($)", "Lean plan ($)", "Projected ($)", "What's in it"], 1):
    c = sm.cell(row, i, h); c.font, c.fill = H_FONT, H_FILL
cats = []
for b in T.BUDGET:
    if b[0] not in cats:
        cats.append(b[0])
desc = {
    "Already paid": "India flights + Buescher cabin",
    "TX – Airport runs": "2 minivan rentals, gas, DFW parking",
    "TX – Day trips": "NASA, Wimberley, Fredericksburg, parks, gas",
    "TX – Food & misc": "Meals out, snacks",
    "Insurance & phones": "Visitor medical insurance, US SIMs",
    "CA – Flights": "United AUS↔CA for 5, bags, airport rides",
    "CA – Lodging": "8 nights incl. taxes & parking",
    "CA – Car & driving": "One-way minivan, CDW, gas, tolls, chains, parking",
    "CA – Activities": "Sequoia pass, Alcatraz, cable cars, aquarium",
    "CA – Food & misc": "Restaurants, groceries, souvenirs",
}
first = row + 1
for cat in cats:
    row += 1
    sm.cell(row, 1, cat).font = BODY
    sm.cell(row, 2, f'=SUMIF(Budget!$A:$A,A{row},Budget!$I:$I)').number_format = USD
    sm.cell(row, 3, f'=SUMIF(Budget!$A:$A,A{row},Budget!$K:$K)').number_format = USD
    sm.cell(row, 4, f'=SUMIF(Budget!$A:$A,A{row},Budget!$O:$O)').number_format = USD
    sm.cell(row, 5, desc.get(cat, "")).font = SUB
    for c in range(1, 5):
        sm.cell(row, c).border = THIN
tot = row + 1
sm.cell(tot, 1, "TOTAL TRIP COST").font = BOLD
sm.cell(tot, 2, f"=SUM(B{first}:B{row})").number_format = USD
sm.cell(tot, 3, f"=SUM(C{first}:C{row})").number_format = USD
sm.cell(tot, 4, f"=SUM(D{first}:D{row})").number_format = USD
for c in (2, 3, 4):
    sm.cell(tot, c).font = BOLD
sm.cell(tot + 1, 1, "Original budget").font = BODY
b = sm.cell(tot + 1, 2, T.BUDGET_TOTAL); b.font = BLUE; b.number_format = USD; b.fill = YELLOW
sm.cell(tot + 1, 3, f"=B{tot + 1}").number_format = USD
sm.cell(tot + 1, 4, f"=B{tot + 1}").number_format = USD
sm.cell(tot + 2, 1, "Over (+) / under (−) budget").font = BOLD
sm.cell(tot + 2, 2, f"=B{tot}-B{tot + 1}").number_format = USD
sm.cell(tot + 2, 3, f"=C{tot}-C{tot + 1}").number_format = USD
sm.cell(tot + 2, 4, f"=D{tot}-D{tot + 1}").number_format = USD
sm.cell(tot + 3, 1, "California part only (Dec 19–27)").font = BODY
sm.cell(tot + 3, 2, f'=SUMIF(Budget!$A:$A,"CA*",Budget!$I:$I)').number_format = USD
sm.cell(tot + 3, 3, f'=SUMIF(Budget!$A:$A,"CA*",Budget!$K:$K)').number_format = USD
sm.cell(tot + 3, 4, f'=SUMIF(Budget!$A:$A,"CA*",Budget!$O:$O)').number_format = USD
sm.cell(tot + 4, 1, "California cost per person").font = BODY
sm.cell(tot + 4, 2, f"=B{tot + 3}/5").number_format = USD
sm.cell(tot + 4, 3, f"=C{tot + 3}/5").number_format = USD
sm.cell(tot + 4, 4, f"=D{tot + 3}/5").number_format = USD
sm.cell(tot + 5, 1, "Paid so far (from Google Sheet Bookings tab)").font = BOLD
sm.cell(tot + 5, 4, f"=Budget!N{BUDGET_TOTAL_ROW}").number_format = USD

n = tot + 7
sm.cell(n, 1, "Key findings").font = BOLD
notes = [
    "Real prices are well above the sheet's $7,670. Main drivers: peak holiday fares, California gas at $6.18/gal (near record), and 5 adults needing a minivan.",
    "Lean plan: Basic Economy fares, a 7-seat SUV instead of a minivan, card-covered CDW, fewer meals out, and skipping the aquarium.",
    "Sequoia NP charges $100 per non-US resident from 2026. Buy one $250 nonresident annual pass for the car (confirm at the gate).",
    "Camp Nelson cabin is booked. Dec 24 means ≈ 5.5 h of driving to see General Sherman. Leave early and check chain status first.",
    "Strong El Niño this winter: carry chains, check Caltrans QuickMap & NPS road status every morning.",
    "Visitor medical insurance for the 3 visitors is the most important 'hidden' cost if not already purchased.",
]
for i, t in enumerate(notes):
    c = sm.cell(n + 1 + i, 1, "• " + t); c.font = BODY
    sm.merge_cells(start_row=n + 1 + i, start_column=1, end_row=n + 1 + i, end_column=5)
    c.alignment = WRAP
    sm.row_dimensions[n + 1 + i].height = 28

# ------------------------------------------------------------------ Itinerary
it = wb.create_sheet("Itinerary", 1)
header(it, 1, ["Date", "Day", "Where", "Plan", "Drive", "Overnight", "Tips / warnings", "Links", "Route"],
       [12, 6, 22, 70, 18, 28, 50, 40, 16])
ROUTES = json.load(open(os.path.join(HERE, "routes.json")))
from build_site import gmaps
for i, d in enumerate(T.DAYS, start=2):
    date = dt.date.fromisoformat(d["date"])
    it.cell(i, 1, date).number_format = "ddd mmm d"
    it.cell(i, 2, f"Day {i - 1}")
    it.cell(i, 3, f'{d["place"]}\n{d["title"]}')
    it.cell(i, 4, "\n".join("• " + p for p in d["plan"]))
    it.cell(i, 5, d["drive"])
    it.cell(i, 6, d["stay"])
    it.cell(i, 7, "\n".join("• " + t for t in d["tips"]))
    lk = [k for k in d["links"] if k in T.L]
    if lk:
        it.cell(i, 8, "\n".join(T.L[k] for k in lk))
        it.cell(i, 8).hyperlink = T.L[lk[0]]
    if d["date"] in T.ROUTES:
        rr = ROUTES[d["date"]]
        it.cell(i, 5, f'{rr["miles"]} mi · {rr["minutes"] // 60} h {rr["minutes"] % 60:02d} m driving')
        rc = it.cell(i, 9, "Google Maps ↗"); rc.hyperlink = gmaps(T.ROUTES[d["date"]])
    fill = CA_FILL if d["phase"] == "California" else TX_FILL
    for c in range(1, 10):
        cell = it.cell(i, c)
        cell.font = LINK if (c == 8 and lk) or (c == 9 and cell.value) else BODY
        cell.alignment = WRAP
        cell.fill = fill
        cell.border = THIN
    it.cell(i, 3).font = BOLD
    it.row_dimensions[i].height = max(45, 15 * max(len(d["plan"]), len(d["tips"]), len(lk), 2))

# ------------------------------------------------------------------ Bookings (edited by the website or by hand)
from openpyxl.worksheet.datavalidation import DataValidation
bk = wb.create_sheet(T.BOOKINGS_TAB, 1)
header(bk, 1, ["ID", "Category", "Item", "Dates", "In plan?", "Planned ($)", "Booked?", "Actual paid ($)", "Paid by", "Notes"],
       [30, 20, 52, 18, 9, 12, 10, 14, 14, 40])
yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
bk.add_data_validation(yn)
for i, row_ in enumerate(BOOK_ROWS, start=2):
    for c, v in enumerate(row_, start=1):
        cell = bk.cell(i, c, v); cell.font = BODY; cell.alignment = WRAP; cell.border = THIN
    for c in (5, 7, 8, 9, 10):
        bk.cell(i, c).fill = YELLOW
    bk.cell(i, 6).number_format = bk.cell(i, 8).number_format = USD2
    yn.add(f"G{i}")
bk.cell(len(BOOK_ROWS) + 3, 2, "The website reads and writes this tab. Edit only the yellow columns; never change the ID column.").font = SUB

# ------------------------------------------------------------------ Options
op = wb.create_sheet("Booking Options")
header(op, 1, ["What", "Option", "Price (live)", "Details", "Status", "Book"], [26, 48, 20, 52, 12, 10])
for i, (sec, opt, price, det, status, lk) in enumerate(T.OPTIONS, start=2):
    for c, v in enumerate([sec, opt, price, det, status], start=1):
        cell = op.cell(i, c, v); cell.font = BODY; cell.alignment = WRAP; cell.border = THIN
    link(op.cell(i, 6), lk, "Book ↗")
    if status == "Chosen":
        for c in range(1, 6):
            op.cell(i, c).font = BOLD
        op.cell(i, 5).fill = YELLOW

r0 = op.max_row + 2
op.cell(r0, 1, "Visitor health insurance in plain words").font = TITLE
for i, (k, v) in enumerate(T.INSURANCE_PRIMER, start=1):
    op.cell(r0 + i, 1, k).font = BOLD
    c = op.cell(r0 + i, 2, v); c.font = BODY; c.alignment = WRAP
    op.merge_cells(start_row=r0 + i, start_column=2, end_row=r0 + i, end_column=4)
    op.row_dimensions[r0 + i].height = 42

# ------------------------------------------------------------------ All links
al = wb.create_sheet("All Links")
header(al, 1, ["Key", "URL"], [20, 120])
for i, (k, u) in enumerate(T.L.items(), start=2):
    al.cell(i, 1, k).font = BODY
    c = al.cell(i, 2, u); c.hyperlink = u; c.font = LINK

# ------------------------------------------------------------------ Weather & safety
ws = wb.create_sheet("Weather & Safety")
ws["A1"] = "December climate (NOAA 1991–2020 normals) & live alerts"; ws["A1"].font = TITLE
ws["A2"] = T.ALERTS_NOTE; ws["A2"].font = SUB; ws["A2"].alignment = WRAP
ws.merge_cells("A2:E2"); ws.row_dimensions[2].height = 60
for i, h in enumerate(["Place", "Avg high °F", "Avg low °F", "Wet days in Dec", "Snow"], 1):
    c = ws.cell(4, i, h); c.font, c.fill = H_FONT, H_FILL
for i, row_ in enumerate(T.CLIMATE, start=5):
    for c, v in enumerate(row_, start=1):
        ws.cell(i, c, v).font = BODY
for col, w in zip("ABCDE", [36, 12, 12, 15, 44]):
    ws.column_dimensions[col].width = w
k = 5 + len(T.CLIMATE) + 1
ws.cell(k, 1, "Check every morning").font = BOLD
for j, (label, key) in enumerate([("Caltrans QuickMap (chain controls, closures)", "quickmap"),
                                   ("Caltrans chain-control rules R1/R2/R3", "chains"),
                                   ("Sequoia NP road conditions (559-565-3341)", "seki_roads"),
                                   ("Sequoia NP current conditions & alerts", "seki_conditions"),
                                   ("NWS forecasts & alerts", "nws"),
                                   ("Enterprise chain policy", "enterprise_chains")], start=1):
    ws.cell(k + j, 1, label).font = BODY
    link(ws.cell(k + j, 2), key)

# ------------------------------------------------------------------ Photo credits
cr = wb.create_sheet("Photo Credits")
header(cr, 1, ["Photo", "File", "Author", "License", "Source"], [18, 50, 30, 14, 12])
credits = json.load(open(os.path.join(HERE, "..", "img", "credits.json")))
for i, (key, v) in enumerate(sorted(credits.items()), start=2):
    for c, val in enumerate([key, v["title"], v["artist"], v["license"]], start=1):
        cr.cell(i, c, val).font = BODY
    c = cr.cell(i, 5, "Commons ↗"); c.hyperlink = v["page"]; c.font = LINK

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = sheet.title in ("All Links",)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("wrote", os.path.abspath(OUT))
